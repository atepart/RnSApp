import tempfile
import unittest
from pathlib import Path

import openpyxl

from domain.constants import DataTableColumns, ParamTableColumns
from domain.models import InitialDataItem, InitialDataItemList
from domain.utils import calculate_rn_sqrt, inverse_diameter_linear_fit
from infrastructure.repository_memory import InMemoryCellRepository
from infrastructure.xlsx_io import WEIGHT_HEADER, XlsxCellIO


class XlsxWeightedFitTests(unittest.TestCase):
    @staticmethod
    def _repo_with_cell() -> InMemoryCellRepository:
        repo = InMemoryCellRepository()
        rows = [
            (1.0, 8.0, True),
            (2.0, 4.0, True),
            (4.0, 2.0, True),
            (8.0, 1.0, True),
            (0.0, 10.0, True),
            (16.0, None, True),
        ]
        initial_data = InitialDataItemList()
        for row, (diameter, resistance, selected) in enumerate(rows):
            values = {
                DataTableColumns.NUMBER: row + 1,
                DataTableColumns.NAME: f"Sample {row + 1}",
                DataTableColumns.SELECT: selected,
                DataTableColumns.DIAMETER: diameter,
                DataTableColumns.RESISTANCE: resistance,
            }
            initial_data.extend(
                InitialDataItem(value=value, row=row, col=column.index) for column, value in values.items()
            )
        repo.update_or_create_item(
            cell=1,
            name="Weighted test",
            diameter_list=[row[0] for row in rows[:4]],
            rn_sqrt_list=[1 / (row[1] ** 0.5) for row in rows[:4]],
            slope=0.1,
            intercept=0.2,
            drift=-2.0,
            rns=10.0,
            drift_error=0.0,
            rns_error=0.0,
            initial_data=initial_data,
        )
        return repo

    def test_export_contains_weighted_formulas_and_fit_series(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            output = Path(tmp_dir) / "weighted.xlsx"
            XlsxCellIO().save(str(output), [("Cell №1", "Weighted test", "")], self._repo_with_cell())
            workbook = openpyxl.load_workbook(output, data_only=False)
            sheet = next(ws for ws in workbook.worksheets if ws.title != "Cells data")

            headers = {cell.value: cell.column for cell in sheet[1] if cell.value is not None}
            weight_col = headers[WEIGHT_HEADER]
            assert weight_col == headers[DataTableColumns.RN_SQRT.slug] + 1
            assert "1/D2" in sheet.cell(row=2, column=weight_col).value
            assert "ISNUMBER(J2)" in sheet.cell(row=2, column=weight_col).value
            assert "D6>0" in sheet.cell(row=6, column=weight_col).value
            assert "ISNUMBER(J7)" in sheet.cell(row=7, column=weight_col).value

            slope_formula = sheet.cell(row=2, column=headers[ParamTableColumns.SLOPE.name]).value
            intercept_formula = sheet.cell(row=2, column=headers[ParamTableColumns.INTERCEPT.name]).value
            wls_helper_cols = {
                sheet.cell(row=1, column=column).value: column
                for column in range(1, sheet.max_column + 1)
                if str(sheet.cell(row=1, column=column).value).startswith("_WLS ")
            }
            assert set(wls_helper_cols) == {
                "_WLS sum w",
                "_WLS sum wD",
                "_WLS sum wY",
                "_WLS sum wD2",
                "_WLS sum wDY",
            }
            for column in wls_helper_cols.values():
                assert "SUMPRODUCT" in sheet.cell(row=2, column=column).value
            assert sheet.cell(row=2, column=wls_helper_cols["_WLS sum w"]).coordinate in slope_formula
            assert sheet.cell(row=2, column=wls_helper_cols["_WLS sum wDY"]).coordinate in slope_formula
            assert "SLOPE(" not in slope_formula
            assert sheet.cell(row=2, column=wls_helper_cols["_WLS sum wY"]).coordinate in intercept_formula

            assert len(sheet._charts) == 1
            assert len(sheet._charts[0].series) == 2
            assert sheet._charts[0].series[0].trendline is None
            hidden_headers = {
                sheet.cell(row=1, column=column).value
                for column in range(1, sheet.max_column + 1)
                if sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].hidden
            }
            assert hidden_headers == {
                "_Weighted fit X",
                "_Weighted fit Y",
                "_WLS sum w",
                "_WLS sum wD",
                "_WLS sum wY",
                "_WLS sum wD2",
                "_WLS sum wDY",
            }

            loaded_items, load_errors = XlsxCellIO().load(str(output))
            assert load_errors == []
            expected_slope, expected_intercept = inverse_diameter_linear_fit(
                [1.0, 2.0, 4.0, 8.0],
                [calculate_rn_sqrt(resistance, 0.0) for resistance in (8.0, 4.0, 2.0, 1.0)],
            )
            self.assertAlmostEqual(loaded_items[0]["slope"], expected_slope, places=12)
            self.assertAlmostEqual(loaded_items[0]["intercept"], expected_intercept, places=12)


if __name__ == "__main__":
    unittest.main()
